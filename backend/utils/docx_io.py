from __future__ import annotations

from io import BytesIO
from typing import List

from docx import Document


def read_docx(file_bytes: bytes) -> List[str]:
    doc = Document(BytesIO(file_bytes))
    return [p.text for p in doc.paragraphs if p.text and p.text.strip()]


def write_docx(paragraphs: List[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
